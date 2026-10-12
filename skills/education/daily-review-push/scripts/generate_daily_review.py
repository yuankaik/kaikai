#!/usr/bin/env python3
"""
佳乐三科每日复习汇总生成脚本 — Word .docx 版本。
由 cron 在 15:00 自动调用，生成桌面 Word 文件。

读取：
  - 语文背诵进度.txt (day, round)
  - 语文背诵100天计划.json
  - 每日复习追踪.json (math + english Ebbinghaus items)

写入：
  - C:/Users/admin/Desktop/佳乐每日复习/YYYY-MM-DD.docx
  - 更新后回写 语文背诵进度.txt + 每日复习追踪.json
"""
import json, os, sys
from datetime import date, timedelta, datetime

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("[ERROR] python-docx not installed. Run: pip install python-docx", file=sys.stderr)
    sys.exit(1)

# === PATHS ===
base = "/mnt/c/Users/admin/Desktop/shanghai-exams"
chinese_progress_path = os.path.join(base, "语文背诵进度.txt")
chinese_plan_path = os.path.join(base, "语文背诵100天计划.json")
review_json_path = os.path.join(base, "每日复习追踪.json")
output_dir = "/mnt/c/Users/admin/Desktop/佳乐每日复习"
os.makedirs(output_dir, exist_ok=True)

today = date.today()
today_str = today.isoformat()

# Idempotency guard: skip if today's file already exists
output_path = os.path.join(output_dir, f"{today_str}.docx")
if os.path.exists(output_path):
    print(f"[SKIP] 今日文件已存在: {output_path}")
    print(f"       如需重新生成，请先删除该文件后重试。")
    sys.exit(0)
ebbinghaus_intervals = [1, 2, 4, 8, 15]
grade_map = {1: "一", 2: "二", 3: "三", 4: "四", 5: "五"}

# === Helper: set font for a run or paragraph ===
def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None, italic=False):
    """Set font properties on a run element."""
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_paragraph_with_font(doc, text, size=11, bold=False, color=None, 
                             alignment=None, space_after=6, space_before=0,
                             first_line_indent=None, left_indent=None):
    """Add a paragraph with unified font settings."""
    p = doc.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if left_indent is not None:
        pf.left_indent = left_indent
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent
    
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p

def add_mixed_paragraph(doc, parts, space_after=6, space_before=0,
                        left_indent=None, first_line_indent=None):
    """Add paragraph with mixed formatting. parts = [(text, size, bold, color, italic), ...]"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if left_indent is not None:
        pf.left_indent = left_indent
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent
    
    for part in parts:
        text, size, bold, color = part[0], part[1], part[2], part[3]
        italic = part[4] if len(part) > 4 else False
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p

# =============================================
# 1. CHINESE (语文背诵)
# =============================================
print("[语文] 读取进度...")
with open(chinese_progress_path, "r", encoding="utf-8") as f:
    lines = f.read().strip().split("\n")
day_num = int(lines[0].strip())
round_num = int(lines[1].strip())
print(f"[语文] 当前进度: Day{day_num} / 第{round_num}轮")

with open(chinese_plan_path, "r", encoding="utf-8") as f:
    plan = json.load(f)

today_items = []
for entry in plan:
    if entry["day"] == day_num:
        today_items = entry["items"]
        break

print(f"[语文] Day{day_num} 共 {len(today_items)} 项")

chinese_lines = []
for item in today_items:
    chinese_lines.append({
        "category": item["category"],
        "grade": item["grade"],
        "title": item["title"],
        "text": item["text"],
        "notes": item.get("notes", ""),
        "author": item.get("author", ""),
        "dynasty": item.get("dynasty", "")
    })

# Update progress
new_day = day_num + 1
new_round = round_num
if new_day > 100:
    new_day = 1
    new_round = round_num + 1

with open(chinese_progress_path, "w", encoding="utf-8") as f:
    f.write(f"{new_day}\n{new_round}\n")
print(f"[语文] 进度已更新: Day{new_day} / 第{new_round}轮")

# =============================================
# 2. MATH & ENGLISH (艾宾浩斯)
# =============================================
print("[数学+英语] 读取 review JSON...")
with open(review_json_path, "r", encoding="utf-8") as f:
    review_data = json.load(f)

items = review_data["items"]

def process_subject(items, subject, max_items, today_str, ebbinghaus_intervals, today):
    """Filter, deduplicate, sort, and update Ebbinghaus items for a subject."""
    due = [it for it in items if it.get("subject") == subject
           and it["next_review"] <= today_str and it["reviews"] < 5]

    # Deduplicate by content
    seen = set()
    unique = []
    for it in due:
        key = it["content"]
        if key not in seen:
            seen.add(key)
            unique.append(it)

    unique.sort(key=lambda x: (x["next_review"], x["id"]))
    selected = unique[:max_items]
    overflow = unique[max_items:]

    print(f"[{subject}] 到期 {len(due)} 条, 去重后 {len(unique)} 条, 今日取 {len(selected)} 条")

    for it in selected:
        it["reviews"] += 1
        new_reviews = it["reviews"]
        if new_reviews < 5:
            delta = ebbinghaus_intervals[new_reviews]
            it["next_review"] = (today + timedelta(days=delta)).isoformat()
        else:
            it["next_review"] = "2999-12-31"  # mastered
        it.setdefault("review_dates", []).append(today_str)

    for it in overflow:
        it["next_review"] = (today + timedelta(days=1)).isoformat()

    # Deduplicated output with context
    seen_out = set()
    output = []
    for it in selected:
        key = it["content"]
        if key not in seen_out:
            seen_out.add(key)
            output.append({
                "content": it["content"],
                "meaning": it.get("meaning", ""),
                "context": it.get("context", ""),
                "reviews": it["reviews"]
            })
    return output

math_output = process_subject(items, "math", 10, today_str, ebbinghaus_intervals, today)
eng_output = process_subject(items, "english", 200, today_str, ebbinghaus_intervals, today)

# Write back JSON
with open(review_json_path, "w", encoding="utf-8") as f:
    json.dump(review_data, f, ensure_ascii=False, indent=2)
print(f"[JSON] 已写回 review JSON ({len(items)} items)")

# =============================================
# 3. GENERATE WORD .docx
# =============================================
print("[文档] 生成 Word .docx...")

doc = Document()

# --- Default style ---
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# Set narrow margins
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# --- Title ---
month_cn = f"{today.month}月{today.day}日"
weekday_cn = ["一", "二", "三", "四", "五", "六", "日"][today.weekday()]

GRAY = RGBColor(0x80, 0x80, 0x80)

add_paragraph_with_font(doc, "佳乐每日复习", size=22, bold=True,
                        alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_paragraph_with_font(doc, f"{today.year}年{month_cn} 星期{weekday_cn}",
                        size=14, color=GRAY,
                        alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

# Divider line
p = doc.add_paragraph()
pf = p.paragraph_format
pf.space_after = Pt(8)
pf.space_before = Pt(0)
run = p.add_run("━" * 40)
set_run_font(run, size=8, color=GRAY)

# =============================================
# CHINESE SECTION
# =============================================
add_paragraph_with_font(doc, f"📖 语文背诵 Day{day_num}（第{round_num}轮）",
                        size=16, bold=True, space_after=4, space_before=8)

for i, item in enumerate(chinese_lines, 1):
    grade_cn = grade_map.get(item["grade"], str(item["grade"]))
    
    # Title line: "[序号]. [分类] 标题（X年级）"
    header_text = f"{i}. [{item['category']}] {item['title']}（{grade_cn}年级）"
    add_paragraph_with_font(doc, header_text, size=11, bold=True, space_after=1)
    
    # Body text with left indent
    text_display = item["text"]
    add_paragraph_with_font(doc, text_display, size=11, 
                            left_indent=Inches(0.3), space_after=1)
    
    # Author/notes - gray small
    note_parts = []
    if item.get("dynasty"):
        note_parts.append(f"〔{item['dynasty']}〕")
    if item.get("author"):
        note_parts.append(item["author"])
    if item.get("notes"):
        note_parts.append(item["notes"])
    if note_parts:
        add_paragraph_with_font(doc, "  ".join(note_parts), size=10, color=GRAY,
                                left_indent=Inches(0.3), space_after=6)

# Divider
p = doc.add_paragraph()
run = p.add_run("━" * 40)
set_run_font(run, size=8, color=GRAY)

# =============================================
# MATH SECTION
# =============================================
add_paragraph_with_font(doc, "📐 数学错题复习" + (f"（{len(math_output)}条）" if math_output else ""),
                        size=16, bold=True, space_after=4, space_before=8)

if math_output:
    for i, item in enumerate(math_output, 1):
        line_text = f"{i}. {item['content']}"
        if item["meaning"]:
            line_text += f" — {item['meaning']}"
        add_paragraph_with_font(doc, line_text, size=11, space_after=1)
        # Context in gray
        if item.get("context") and item["context"] != "来自Day3回炉版":
            add_paragraph_with_font(doc, f"    {item['context']}", size=9, color=GRAY,
                                    space_after=4)
else:
    add_paragraph_with_font(doc, "今日无到期错题 ✅", size=11, color=GRAY, space_after=6)

# Divider
p = doc.add_paragraph()
run = p.add_run("━" * 40)
set_run_font(run, size=8, color=GRAY)

# =============================================
# ENGLISH SECTION
# =============================================
add_paragraph_with_font(doc, "🔤 英语单词复习" + (f"（{len(eng_output)}词）" if eng_output else ""),
                        size=16, bold=True, space_after=4, space_before=8)

if eng_output:
    for i, item in enumerate(eng_output, 1):
        # Mixed paragraph: number + word (bold 12pt) + dash + meaning
        add_mixed_paragraph(doc, [
            (f"{i}. ", 11, False, None),
            (item["content"], 12, True, None),
            (f" — {item['meaning']}", 11, False, None),
        ], space_after=2)
        # Context in gray
        if item.get("context"):
            add_paragraph_with_font(doc, f"    {item['context']}", size=9, color=GRAY,
                                    space_after=4)
else:
    add_paragraph_with_font(doc, "今日无到期单词 ✅", size=11, color=GRAY, space_after=6)

# Footer
p = doc.add_paragraph()
run = p.add_run("━" * 40)
set_run_font(run, size=8, color=GRAY)

now = datetime.now()
time_str = now.strftime("%H:%M")
add_paragraph_with_font(doc, f"📌 生成时间：{time_str}  |  明日继续加油！💪",
                        size=9, color=GRAY, 
                        alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=0,
                        space_before=8)

# =============================================
# SAVE
# =============================================
filename = f"{today_str}.docx"
output_path = os.path.join(output_dir, filename)
doc.save(output_path)

file_size = os.path.getsize(output_path)
print("=" * 50)
print(f"✅ 文件已生成: {output_path}")
print(f"📏 文件大小: {file_size:,} bytes ({file_size/1024:.1f} KB)")
print(f"📖 语文: Day{day_num} 第{round_num}轮 {len(chinese_lines)}项")
print(f"📐 数学: {len(math_output)}条错题")
print(f"🔤 英语: {len(eng_output)}个单词")
print("=" * 50)
print("=== DONE ===")
