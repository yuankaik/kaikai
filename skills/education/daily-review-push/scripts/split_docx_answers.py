#!/usr/bin/env python3
"""
批量拆分练习卷：学生卷（短文+题目） / 家长答案（答案+批改提示+重写任务）

用法：
  python3 split_docx_answers.py <输入目录> [输出目录]

默认输出到输入目录下的 学生卷/ 和 家长答案/ 子目录。
"""

import sys, os, re
from docx import Document


def copy_paragraph(src_p, dst_doc):
    """复制段落，保留粗体/斜体/字号"""
    p = dst_doc.add_paragraph()
    p.alignment = src_p.alignment
    for run in src_p.runs:
        r = p.add_run(run.text)
        r.bold = run.bold
        r.italic = run.italic
        r.underline = run.underline
        if run.font.size:
            r.font.size = run.font.size
        if run.font.name:
            r.font.name = run.font.name
    return p


def is_parent_section_header(text):
    """检测家长批改提示/孩子重写任务区段头（多种格式）"""
    clean = text.replace('**', '').strip()
    patterns = [
        r'^短文[一二]\s*[··]\s*家长批改提示',
        r'^短文[一二]家长批改提示',
        r'^家长批改提示',
        r'^孩子重写任务',
    ]
    for pat in patterns:
        if re.match(pat, clean):
            return True
    if len(clean) < 30 and ('家长批改提示' in clean or '孩子重写任务' in clean):
        return True
    return False


def is_passage_boundary(text):
    """区段退出：新短文开始"""
    if re.match(r'^(阅读短文|短文)[一二二]', text):
        return True
    if re.match(r'^【第[一二]篇】', text):
        return True
    if text.startswith('今日总结') or text.startswith('明天预告') or text.startswith('本日训练要点'):
        return True
    return False


def is_answer_content(text):
    """判断是否为答案/评分内容"""
    if '标准答案' in text:
        return True
    if '依据段' in text and ('：' in text or ':' in text):
        return True
    if '评分点' in text and len(text) < 20:
        return True
    if '常见错因标签' in text or '常见错因' in text:
        return True
    if re.match(r'^[•·]\s*.*\d+\s*分', text):
        return True
    if re.match(r'^[•·]\s*答案以', text):
        return True
    if re.match(r'^[•·]\s*共\d+分', text):
        return True
    return False


def split_docx(filepath):
    """拆分为 (student_doc, parent_doc)"""
    doc = Document(filepath)
    student_doc = Document()
    parent_doc = Document()
    in_parent_section = False

    for p in doc.paragraphs:
        text = p.text.strip()

        # 父卷保留全文
        copy_paragraph(p, parent_doc)

        # 区段切换
        if is_parent_section_header(text):
            in_parent_section = True
            continue

        if is_passage_boundary(text):
            in_parent_section = False

        if in_parent_section:
            continue

        if is_answer_content(text):
            continue

        copy_paragraph(p, student_doc)

    return student_doc, parent_doc


def main():
    src_dir = sys.argv[1] if len(sys.argv) > 1 else '.'
    out_dir = sys.argv[2] if len(sys.argv) > 2 else src_dir

    student_dir = os.path.join(out_dir, '学生卷')
    parent_dir = os.path.join(out_dir, '家长答案')
    os.makedirs(student_dir, exist_ok=True)
    os.makedirs(parent_dir, exist_ok=True)

    files = sorted(f for f in os.listdir(src_dir) if f.endswith('.docx') and not f.startswith('~'))
    clean = 0

    for fname in files:
        fpath = os.path.join(src_dir, fname)
        try:
            sdoc, pdoc = split_docx(fpath)
            base = fname.replace('阅读理解练习', '')
            sdoc.save(os.path.join(student_dir, base.replace('.docx', '练习卷.docx') if '练习' not in base else base))
            pdoc.save(os.path.join(parent_dir, base.replace('.docx', '答案与指导.docx') if '答案' not in base else base))

            stext = '\n'.join([p.text for p in sdoc.paragraphs])
            if '标准答案' not in stext and '家长批改提示' not in stext:
                clean += 1
            else:
                print(f"  ⚠ {fname}: 泄漏检查未通过")
        except Exception as e:
            print(f"  ✗ {fname}: {e}")

    print(f"\n完成: {clean}/{len(files)} 个文件拆分干净")


if __name__ == '__main__':
    main()
