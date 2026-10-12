"""
Generate a merged vocab DOCX with the same format as desktop 单词.docx.

Input: merged words list (current + missing), output path.
Format: groups of 20 words, tab-separated: 序号 \t 单词 \t 中文
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def generate_merged_docx(words, output_path, title='词汇汇总（完整版）',
                          subtitle='', group_size=20, current_count=None):
    """
    Args:
        words: list of {english, chinese}
        output_path: absolute path for .docx
        title: document title
        subtitle: subtitle line
        group_size: words per group (default 20)
        current_count: if set, words before this index are 'already in practice',
                       words after get '← 补缺' marker
    """
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    
    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(subtitle)
        run.font.size = Pt(9)
        run.font.name = '宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    doc.add_paragraph()
    
    group_num = 1
    for i in range(0, len(words), group_size):
        group = words[i:i+group_size]
        
        suffix = ''
        if current_count and i >= current_count:
            suffix = '  ← 补缺'
        
        # Group header
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f'{group_num:02d}{suffix}')
        run.font.size = Pt(11)
        run.font.bold = True
        
        # Column header
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run('序号\t单词\t中文')
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.name = '宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        
        for idx, w in enumerate(group, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(0)
            run = p.add_run(f'{idx}\t{w.get("english","")}\t{w.get("chinese","")}')
            run.font.size = Pt(9)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        
        doc.add_paragraph()
        group_num += 1
    
    doc.save(output_path)
    return len(words), group_num - 1
