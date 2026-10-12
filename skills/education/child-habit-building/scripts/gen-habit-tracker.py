# -*- coding: utf-8 -*-
"""Generate father-son 66-day habit tracker - minimal 2 habits each"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import date, timedelta

start = date(2026, 8, 1)  # CHANGE start date as needed

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.5); s.bottom_margin = Cm(1)
    s.left_margin = Cm(2); s.right_margin = Cm(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('父子 66 天习惯打卡')
run.font.size = Pt(20); run.font.bold = True
run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(f'{start.month}月{start.day}日 → {(start+timedelta(65)).month}月{(start+timedelta(65)).day}日 | 互相督促')
run.font.size = Pt(10)
run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

doc.add_paragraph()

# Dad habits
p = doc.add_paragraph()
run = p.add_run('■ 爸')
run.font.size = Pt(12); run.font.bold = True
run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
p = doc.add_paragraph()
run = p.add_run('① 早上不刷手机  ② 睡前写3件事')
run.font.size = Pt(10)
run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# Child habits  
doc.add_paragraph()
p = doc.add_paragraph()
run = p.add_run('■ 乐')
run.font.size = Pt(12); run.font.bold = True
run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
p = doc.add_paragraph()
run = p.add_run('① 晨读10分钟  ② 晚间背诵20分钟')
run.font.size = Pt(10)
run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

p = doc.add_paragraph()
run = p.add_run('✓做到了  □没做到  周末不放假')
run.font.size = Pt(9)
run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

doc.add_paragraph()

# Table
table = doc.add_table(rows=1, cols=7)
table.style = 'Table Grid'
headers = ['日期', '天数', '爸①', '爸②', '乐①', '乐②', '备注']
widths = [Cm(2.8), Cm(0.9), Cm(1.3), Cm(1.3), Cm(1.3), Cm(1.3), Cm(5)]

for i, (h, w) in enumerate(zip(headers, widths)):
    cell = table.rows[0].cells[i]; cell.width = w
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h); run.font.size = Pt(9); run.font.bold = True
    run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'D9D9D9')
    cell._element.get_or_add_tcPr().append(shading)

for d in range(66):
    cur = start + timedelta(days=d)
    wn = ['一','二','三','四','五','六','日'][cur.weekday()]
    row = table.add_row()
    
    p = row.cells[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'{cur.month}月{cur.day}日 周{wn}')
    run.font.size = Pt(8)
    run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    p = row.cells[1].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(d+1)); run.font.size = Pt(8)
    
    for j in [2,3,4,5]:
        p = row.cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run('☐'); run.font.size = Pt(11)
    
    p = row.cells[6].paragraphs[0]
    note = {0:'←开始',6:'第1周',20:'21天！',41:'42天',65:'66天！'}.get(d,'')
    run = p.add_run(note); run.font.size = Pt(7)

doc.add_paragraph()
p = doc.add_paragraph()
run = p.add_run('漏了不打不骂不补，第二天继续 ✓ 就好。')
run.font.size = Pt(9); run.font.italic = True
run.font.name = '宋体'; run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

doc.save('/mnt/c/Users/admin/Desktop/父子66天习惯打卡.docx')
print('Done')
