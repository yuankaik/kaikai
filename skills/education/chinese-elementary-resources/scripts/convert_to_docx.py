#!/usr/bin/env python3
"""Convert markdown files to printable DOCX with Chinese formatting.
Requires: python-docx (pip install python-docx)
Usage: python3 convert_to_docx.py <source_dir> [output_dir]
"""

import os, re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_font(run, name='宋体', size=11, bold=False, color=None):
    """Set Chinese font with proper east-asia handling."""
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        from lxml import etree
        rFonts = etree.SubElement(rPr, qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, indent=True):
    """Add a normal paragraph with bold support via **text**."""
    if not text.strip():
        return
    p = doc.add_paragraph()
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            set_font(run, '宋体', 11, bold=True)
        else:
            run = p.add_run(part)
            set_font(run, '宋体', 11)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)


def add_heading_styled(doc, text, level):
    """Add heading with Chinese styling."""
    p = doc.add_paragraph()
    sizes = {1: 18, 2: 15, 3: 13, 4: 12}
    size = sizes.get(level, 11)
    run = p.add_run(text.lstrip('# '))
    set_font(run, '黑体', size, bold=True)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_separator(doc):
    """Add a subtle horizontal separator."""
    p = doc.add_paragraph()
    run = p.add_run('─' * 40)
    set_font(run, '宋体', 8, color=(180, 180, 180))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def convert_file(md_path, docx_path):
    """Convert a single markdown file to DOCX."""
    doc = Document()

    # A4 page setup
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        line = line.rstrip()
        if not line:
            continue

        # Headings
        if line.startswith('#### '):
            add_heading_styled(doc, line, 4)
        elif line.startswith('### '):
            add_heading_styled(doc, line, 3)
        elif line.startswith('## '):
            add_heading_styled(doc, line, 2)
        elif line.startswith('# '):
            add_heading_styled(doc, line, 1)
        elif line.strip() == '---':
            add_separator(doc)
        elif line.startswith('> '):
            # Blockquote
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            set_font(run, '楷体', 10.5, color=(100, 100, 100))
            p.paragraph_format.left_indent = Cm(1)
        elif line.startswith('|'):
            # Simple table row
            add_para(doc, line.replace('|', ' | ').strip(' |'), indent=False)
        elif line.startswith('- ') or line.startswith('  - '):
            add_para(doc, '• ' + line.lstrip('- '))
        else:
            add_para(doc, line)

    os.makedirs(os.path.dirname(docx_path) if os.path.dirname(docx_path) else '.', exist_ok=True)
    doc.save(docx_path)
    return True


def batch_convert(source_dir, output_dir=None):
    """Convert all .md files in source_dir to .docx in output_dir."""
    if output_dir is None:
        output_dir = source_dir + '-docx'

    os.makedirs(output_dir, exist_ok=True)

    md_files = []
    for root, dirs, files in os.walk(source_dir):
        for f in files:
            if f.endswith('.md'):
                md_files.append(os.path.join(root, f))

    for i, md_path in enumerate(sorted(md_files)):
        rel = os.path.relpath(md_path, source_dir)
        docx_name = rel.replace('/', '_').replace('\\', '_').replace('.md', '.docx')
        docx_path = os.path.join(output_dir, docx_name)
        try:
            convert_file(md_path, docx_path)
            print(f'[{i+1}/{len(md_files)}] ✓ {rel}')
        except Exception as e:
            print(f'[{i+1}/{len(md_files)}] ✗ {rel}: {e}')

    print(f'\nDone! {len(md_files)} files → {output_dir}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Usage: python3 convert_to_docx.py <source_dir> [output_dir]')
        sys.exit(1)

    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else None
    batch_convert(src, out)
