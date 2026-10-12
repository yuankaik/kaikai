#!/usr/bin/env python3
"""
验证 40 天暑假练习 DOCX 文件质量。
检查项：A4尺寸、字体、字号、答案泄漏、模块完整性、答题空行数、舒尔特布局。
用法：python3 scripts/verify_40days.py
"""
from docx import Document
from docx.oxml.ns import qn
import glob, re, os

BASE = "/mnt/c/Users/admin/Desktop/佳乐2026暑假-精简版"

LEAK_KW = ["答案与解析", "标准答案", "参考答案", "方便家长精准指导"]
MODULES = ["🔢 口算", "📐 数学", "🧠 舒尔特", "📝 语文基础",
           "📖 语文阅读", "🌅 晨读", "🌍 英语阅读", "📐 英语语法"]


def is_q(text):
    return bool(re.match(r'^(第\d+题|题\d+)', text))


def verify_day(day):
    week = (day - 1) // 5 + 1
    week_dirs = glob.glob(f"{BASE}/第{week}周*")
    if not week_dirs:
        return "Week dir not found"
    
    path = f"{week_dirs[0]}/学生卷/Day{day:02d}-全部练习.docx"
    if not os.path.exists(path):
        return "File not found"
    
    doc = Document(path)
    full = "\n".join(p.text for p in doc.paragraphs)
    issues = []
    
    # Check 1: A4 page size (210×297mm ±1mm)
    s = doc.sections[0]
    w_mm = s.page_width / 36000
    h_mm = s.page_height / 36000
    if abs(w_mm - 210) > 1 or abs(h_mm - 297) > 1:
        issues.append(f"Page not A4: {w_mm:.0f}x{h_mm:.0f}mm")
    
    # Check 1b: minimum margin (left >= 15mm for binding)
    lm = s.left_margin / 36000
    if lm < 15:
        issues.append(f"Left margin too small: {lm:.0f}mm")
    
    # Check 1c: font audit (only 宋体 + Arial for Schulte grids)
    bad_fonts = set()
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.name and r.font.name not in ('宋体', 'Arial'):
                bad_fonts.add(r.font.name)
    if bad_fonts:
        issues.append(f"Non-standard fonts: {bad_fonts}")
    
    # Check 1d: font sizes (no <8pt content, no >20pt)
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        for r in p.runs:
            if r.font.size:
                sz = r.font.size / 12700
                if sz < 8:
                    issues.append(f"Tiny font: {sz:.0f}pt at '{p.text[:30]}'")
                    break
                if sz > 20:
                    issues.append(f"Oversized font: {sz:.0f}pt at '{p.text[:30]}'")
                    break
    
    # Check 2: answer leaks
    for kw in LEAK_KW:
        if kw in full:
            issues.append(f"LEAK: '{kw}'")
    
    # Check 3: all modules present
    for mod in MODULES:
        if mod not in full:
            issues.append(f"MISSING: {mod}")
    
    # Check 4: schulte count (expected 2)
    schulte_count = full.count("舒尔特")
    if schulte_count != 2:
        issues.append(f"Schulte appears {schulte_count}x (expected 2)")
    
    # Check 5: reading passages (should be 2)
    passage_count = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if '阅读短文' in t and ('一' in t or '二' in t or '：《' in t or '：' in t):
            if len(t) < 60 and not t.startswith('📖'):
                passage_count += 1
    if passage_count < 2 and day not in [8, 9]:
        issues.append(f"Reading passages: {passage_count} (expected 2)")
    
    # Check 6: question blanks (sampling - check if any question has 0 blanks)
    in_reading = False
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if '语文阅读' in text:
            in_reading = True
            continue
        if in_reading and '晨读' in text:
            in_reading = False
        if not in_reading:
            continue
        if is_q(text):
            body_end = i
            for j in range(i + 1, len(doc.paragraphs)):
                nt = doc.paragraphs[j].text.strip()
                if is_q(nt) or nt.startswith('──'):
                    break
                if nt:
                    body_end = j
            blanks = sum(1 for j in range(body_end + 1, min(body_end + 5, len(doc.paragraphs)))
                        if not doc.paragraphs[j].text.strip())
            if blanks == 0:
                issues.append(f"Q P{i}: 0 blanks")
                break
    
    if issues:
        return "; ".join(issues)
    return "OK"


if __name__ == "__main__":
    ok = 0
    bad = 0
    for day in range(1, 41):
        result = verify_day(day)
        if result == "OK":
            ok += 1
        else:
            bad += 1
            print(f"Day {day:02d}: {result}")
    
    print(f"\n{ok}/40 OK, {bad}/40 issues")
