#!/usr/bin/env python3
"""
通用模板：从转录MD文件夹生成晨读作文Word文档。

用法：
  1. 转录MD放在一个文件夹（如 佳乐晨读作文60篇/）
  2. 修改下方的 SRC 和 OUT 路径
  3. 运行：python3 scripts/essays_to_docx.py

MD格式要求：
  # 作文标题
  > 分类: xxx
  ---
  正文内容...
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os, re

# ===== 配置区 =====
SRC = "/mnt/c/Users/admin/Desktop/佳乐晨读作文60篇"  # MD源文件夹
OUT = "/mnt/c/Users/admin/Desktop/佳乐晨读作文60篇.docx"  # 输出Word
TITLE = "佳乐晨读作文60篇"
SUBTITLE = "小学3-5年级优秀范文精选"
FOOTER = "每日晨读一篇 · 两个月读完"

# ===== 同音字修正表（whisper base模型常见错误）=====
FIXES = [
    ('由如', '犹如'), ('種種有神', '炯炯有神'), ('宇文老师', '语文老师'),
    ('爛成了一鍋舟', '乱成了一锅粥'), ('荷艾', '和蔼'), ('打科睡', '打瞌睡'),
    ('轟躺大笑', '哄堂大笑'), ('熬油', '遨游'), ('槍属', '仓鼠'),
    ('原流流', '圆溜溜'), ('亮金金', '亮晶晶'), ('老师动来动去', '老是动来动去'),
    ('紧身', '谨慎'), ('刀处跑', '到处跑'), ('文起了躲茫茫', '玩起了捉迷藏'),
    ('艰果', '坚果'), ('飞崩', '飞奔'), ('摊吃', '贪吃'), ('不经', '不禁'),
    ('選资', '选自'), ('皂纹', '皱纹'), ('勾合', '沟壑'), ('稀稀舒舒', '稀稀疏疏'),
    ('碗如', '宛如'), ('到印', '倒映'), ('清山', '青山'), ('蘇滴', '苏堤'),
    ('洋柳一一', '杨柳依依'), ('极耻而过', '疾驰而过'), ('见起', '溅起'),
    ('示威威咒', '微微皱'), ('辭起彼符', '此起彼伏'), ('想撤銀銷', '响彻云霄'),
    ('一馬當仙', '一马当先'), ('腳踐有力', '矫健有力'), ('堅定自行', '坚定自信'),
    ('情轉多餘', '晴转多云'), ('上課林想', '上课铃响'), ('課堅', '课间'),
    ('倒大', '捣蛋'), ('眼力', '严厉'), ('做職', '坐直'),
    ('忽如生', '呼噜声'), ('一觉', '一绝'),
]

# ===== 分类体系 =====
CAT_NAMES = {
    '写人': '👤 写人篇', '写事': '🏃 写事篇', '状物': '🐾 状物篇',
    '写景': '🌄 写景篇', '感悟': '💡 感悟篇', '读后感': '📖 读后感篇',
    '说明': '🔍 说明篇', '推荐一本书': '📚 推荐一本书', '其他': '📝 其他'
}
CAT_ORDER = {'写人':0,'写事':1,'状物':2,'写景':3,'感悟':4,'读后感':5,'说明':6,'推荐一本书':7,'其他':8}

def classify(title):
    """根据标题自动分类"""
    keywords = {
        '写人': ['老师', '爷爷', '奶奶', '妈妈', '爸爸', '同学', '同桌', '朋友'],
        '写事': ['扳手腕', '活动', '运动会', '比赛', '足球', '拔河', '记一次'],
        '状物': ['心爱之物', '白鹭', '小猫', '小狗', '仓鼠'],
        '写景': ['西湖', '搭石', '中秋节', '春天', '秋天'],
        '感悟': ['温暖', '雨中', '风雨', '背影', '那一刻'],
        '读后感': ['读后感'],
        '说明': ['什么比'],
        '推荐一本书': ['推荐一本书'],
    }
    for cat, kws in keywords.items():
        if any(kw in title for kw in kws):
            return cat
    return '其他'

def clean_text(text):
    for old, new in FIXES:
        text = text.replace(old, new)
    return text.strip()

def make_docx(essays, out_path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '宋体'; style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for sec in doc.sections:
        sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)

    # Cover
    for _ in range(6): doc.add_paragraph()
    for text, size, color, bold in [
        (TITLE, 28, None, True), (SUBTITLE, 16, RGBColor(0x66,0x66,0x66), False),
        ("", 12, None, False), (FOOTER, 12, RGBColor(0x99,0x99,0x99), False),
    ]:
        if not text: doc.add_paragraph(); continue
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        if color: r.font.color.rgb = color

    # Category summary page
    doc.add_page_break()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("内容分类"); r.bold = True; r.font.size = Pt(18)
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    doc.add_paragraph()
    from collections import Counter
    for cat in ['写人','写事','状物','写景','感悟','读后感','说明','推荐一本书','其他']:
        cnt = sum(1 for e in essays if e['cat'] == cat)
        if cnt:
            p = doc.add_paragraph()
            r = p.add_run(f"{CAT_NAMES.get(cat, cat)} × {cnt}篇")
            r.font.size = Pt(12); r.font.name = '微软雅黑'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # Essays
    current_cat = ""
    for i, essay in enumerate(essays):
        if essay['cat'] != current_cat:
            current_cat = essay['cat']
            doc.add_page_break()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(CAT_NAMES.get(current_cat, current_cat))
            r.bold = True; r.font.size = Pt(20)
            r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            doc.add_paragraph()

        # Number + title
        for text, size, color in [(f"第{i+1}篇", 10, RGBColor(0xAA,0xAA,0xAA)), (essay['title'], 16, None)]:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text); r.font.size = Pt(size)
            r.font.name = '微软雅黑' if size > 12 else '宋体'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑' if size > 12 else '宋体')
            if color: r.font.color.rgb = color
            if size > 12: r.bold = True
        doc.add_paragraph()

        # Body
        for para in essay['text'].split('\n'):
            para = para.strip()
            if not para: continue
            # Split long paragraphs
            chunks = re.split(r'([。！？])', para)
            buf = ''
            for j, chunk in enumerate(chunks):
                buf += chunk
                if j % 2 == 1 and len(buf) > 40:
                    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0.7)
                    p.paragraph_format.line_spacing = 1.8
                    r = p.add_run(buf); r.font.size = Pt(12)
                    r.font.name = '宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                    buf = ''
            if buf:
                p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0.7)
                p.paragraph_format.line_spacing = 1.8
                r = p.add_run(buf); r.font.size = Pt(12)
                r.font.name = '宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

        # Source note
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("——选自《小学作文范文》"); r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xBB,0xBB,0xBB)
        r.font.name = '宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    doc.save(out_path)
    return len(essays)

# ===== Main =====
if __name__ == '__main__':
    essays = []
    for f in sorted(os.listdir(SRC)):
        if not f.endswith('.md'): continue
        with open(os.path.join(SRC, f), 'r', encoding='utf-8') as fh:
            content = fh.read()
        lines = content.split('\n')
        title = ""; text_lines = []; in_text = False
        for line in lines:
            if line.startswith('# ') and not title: title = line[2:].strip()
            elif line.strip() == '---' and title: in_text = True
            elif in_text and line.strip(): text_lines.append(line.strip())
        if title and text_lines:
            text = clean_text('\n'.join(text_lines))
            essays.append({'title': title, 'text': text, 'cat': classify(title), 'chars': len(text)})

    essays.sort(key=lambda e: (CAT_ORDER.get(e['cat'], 99), e['title']))
    count = make_docx(essays, OUT)
    print(f"✅ {count}篇 → {OUT} ({os.path.getsize(OUT)//1024}KB)")
    for cat in ['写人','写事','状物','写景','感悟','读后感','说明','推荐一本书','其他']:
        cnt = sum(1 for e in essays if e['cat'] == cat)
        if cnt: print(f"   {CAT_NAMES.get(cat, cat)}: {cnt}篇")
