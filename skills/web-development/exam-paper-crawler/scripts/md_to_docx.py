#!/usr/bin/env python3
"""Convert markdown files to printable DOCX format.
Usage: python3 scripts/md_to_docx.py <input.md> <output.docx>
       python3 scripts/md_to_docx.py --dir <directory>
"""
import os, re, sys, argparse
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

def set_font(run, name='宋体', size=11, bold=False):
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = __import__('lxml.etree', fromlist=['etree']).SubElement(rPr, qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)

def add_para(doc, text, indent=True):
    if not text.strip(): return
    p = doc.add_paragraph()
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            set_font(run, '宋体', 11, bold=True)
        else:
            run = p.add_run(part)
            set_font(run, '宋体', 11)
    if indent: p.paragraph_format.first_line_indent = Cm(0.74)

def add_h(doc, text, level):
    p = doc.add_paragraph()
    sizes = {1: 18, 2: 15, 3: 13, 4: 12}
    run = p.add_run(text.lstrip('# '))
    set_font(run, '黑体', sizes.get(level, 11), bold=True)

def convert(md_path, docx_path):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.5)
    with open(md_path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.rstrip()
            if not line: continue
            if line.startswith('#### '): add_h(doc, line, 4)
            elif line.startswith('### '): add_h(doc, line, 3)
            elif line.startswith('## '): add_h(doc, line, 2)
            elif line.startswith('# '): add_h(doc, line, 1)
            elif line.startswith('|'): add_para(doc, line.replace('|', ' | ').strip(' |'))
            elif line.startswith('- ') or line.startswith('  - '): add_para(doc, '  '+line.lstrip('- '))
            elif line.strip() == '---': pass
            else: add_para(doc, line)
    os.makedirs(os.path.dirname(docx_path) or '.', exist_ok=True)
    doc.save(docx_path)
    return os.path.getsize(docx_path)

def main():
    parser = argparse.ArgumentParser(description='MD to DOCX converter')
    parser.add_argument('input', nargs='?', help='Input .md file')
    parser.add_argument('output', nargs='?', help='Output .docx file')
    parser.add_argument('--dir', help='Convert all .md files in directory')
    args = parser.parse_args()

    if args.dir:
        total = 0
        for f in sorted(os.listdir(args.dir)):
            if f.endswith('.md'):
                src = os.path.join(args.dir, f)
                dst = os.path.join(args.dir, f.replace('.md', '.docx'))
                sz = convert(src, dst)
                print(f'  {f} -> {sz//1024}KB')
                total += 1
        print(f'Done: {total} files')
    elif args.input and args.output:
        sz = convert(args.input, args.output)
        print(f'{args.input} -> {args.output} ({sz//1024}KB)')
    else:
        parser.print_help()

if __name__ == '__main__':
    main()
